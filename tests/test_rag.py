"""RAG 兼容行为的离线测试；不调用付费模型或真实存储。"""

from collections.abc import AsyncIterator
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage, AIMessageChunk
from pypdf import PdfWriter

from app.core.config import Settings
from app.main import app
from app.rag.models import FileInfo
from app.rag.parser import MAX_TEXT_LENGTH, TRUNCATION_SUFFIX, parse_file
from app.rag.retriever import RagRetriever
from app.rag.service import FileService
from app.rag.splitter import split_overlap_paragraphs
from app.runtime.file_chat import FileChatRunner


class MemoryFiles:
    def __init__(self) -> None:
        self.items: dict[str, FileInfo] = {}

    async def create(self, info: FileInfo) -> None:
        self.items[info.file_id] = info

    async def update(self, info: FileInfo) -> None:
        self.items[info.file_id] = info

    async def get(self, file_id: str) -> FileInfo | None:
        return self.items.get(file_id)

    async def list_all(self) -> list[FileInfo]:
        return list(self.items.values())

    async def delete(self, file_id: str) -> None:
        self.items.pop(file_id)


class MemoryObjects:
    def __init__(self) -> None:
        self.names: list[str] = []
        self.deleted: list[str] = []

    async def put(self, name: str, content: bytes, content_type: str) -> str:
        self.names.append(name)
        return f"http://minio/bucket/{name}"

    async def delete(self, name: str) -> None:
        self.deleted.append(name)


class MemoryIndex:
    def __init__(self) -> None:
        self.replaced: list[tuple[str, list[str]]] = []
        self.deleted: list[str] = []
        self.searched: list[str] = []

    async def replace(self, file_id: str, chunks: list[str]) -> None:
        self.replaced.append((file_id, chunks))

    async def delete(self, file_id: str) -> None:
        self.deleted.append(file_id)

    async def search(self, file_id: str, query: str, top_k: int) -> list[tuple[str, str]]:
        assert top_k == 5
        self.searched.append(file_id)
        return [("a", "第一片段"), ("b", "第二片段")]


class QueryModel:
    def __init__(self) -> None:
        self.calls = 0

    async def ainvoke(self, messages: list) -> AIMessage:
        self.calls += 1
        return AIMessage(
            content="压缩的问题" if self.calls == 1 else '["问法一", "问法二", "问法三"]'
        )


class ChatModel:
    def __init__(self) -> None:
        self.calls = 0

    def bind_tools(self, tools: list) -> "ChatModel":
        return self

    async def astream(self, messages: list) -> AsyncIterator[AIMessageChunk]:
        self.calls += 1
        if self.calls == 1:
            yield AIMessageChunk(
                content="",
                tool_calls=[
                    {
                        "name": "loadContent",
                        "args": {"fileId": "other-file", "question": "问"},
                        "id": "call-1",
                    }
                ],
            )
        else:
            assert any(
                getattr(message, "content", "") == "文件ID与当前会话不匹配"
                for message in messages
            )
            yield AIMessageChunk(content="文件内容不足，无法回答。")


class MemorySessions:
    def __init__(self) -> None:
        self.saved: dict = {}

    async def recent_file(self, conversation_id: str, file_id: str, limit: int):
        return []

    async def create(self, conversation_id: str, query: str, **kwargs) -> int:
        self.saved.update(kwargs)
        return 1

    async def finish(self, record_id: int, **kwargs) -> None:
        self.saved.update(kwargs)


def make_service(
    index: MemoryIndex | None = None,
) -> tuple[FileService, MemoryFiles, MemoryObjects]:
    files = MemoryFiles()
    objects = MemoryObjects()
    settings = Settings(_env_file=None)
    return FileService(files, objects, index, None, settings), files, objects


def test_java_splitter_and_display_truncation() -> None:
    chunks = split_overlap_paragraphs("a" * 500 + "\n" + "b" * 10)
    assert chunks == ["a" * 500, "a" * 50 + "b" * 10]
    result = parse_file(("中" * (MAX_TEXT_LENGTH + 1)).encode(), "txt")
    assert len(result.full_text) == MAX_TEXT_LENGTH + 1
    assert result.truncated_text == "中" * MAX_TEXT_LENGTH + TRUNCATION_SUFFIX
    with pytest.raises(ValueError, match="暂不支持 .doc"):
        parse_file(b"x", "doc")


def test_docx_and_pdf_parser_adapters_load_real_file_bytes() -> None:
    document = Document()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    docx_bytes = BytesIO()
    document.save(docx_bytes)
    assert parse_file(docx_bytes.getvalue(), "docx").full_text == "第一段\n第二段"

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    pdf_bytes = BytesIO()
    writer.write(pdf_bytes)
    assert parse_file(pdf_bytes.getvalue(), "pdf").full_text == ""


@pytest.mark.asyncio
async def test_upload_uses_5000_character_threshold_and_legacy_metadata() -> None:
    index = MemoryIndex()
    service, files, objects = make_service(index)
    small = await service.upload("small.txt", b"a" * 4999, "text/plain")
    large = await service.upload("large.txt", b"b" * 5000, "text/plain")
    assert small.embed == 0 and large.embed == 1
    assert small.status == large.status == "SUCCESS"
    assert len(index.replaced) == 1
    assert index.replaced[0][0] == large.file_id
    assert objects.names[0].startswith("file-")
    assert small.to_legacy_dict()["fileId"] == small.file_id
    assert files.items[large.file_id].embed == 1
    await service.delete(large.file_id)
    assert index.deleted == [large.file_id]
    assert large.file_id not in files.items


@pytest.mark.asyncio
async def test_retriever_uses_four_queries_file_filter_and_deduplicates() -> None:
    index = MemoryIndex()
    model = QueryModel()
    results = await RagRetriever(model, index).retrieve("only-this-file", "原问题")
    assert results == ["第一片段", "第二片段"]
    assert model.calls == 2
    assert index.searched == ["only-this-file"] * 4


@pytest.mark.asyncio
async def test_index_failure_keeps_java_display_text_fallback() -> None:
    class BrokenIndex(MemoryIndex):
        async def replace(self, file_id: str, chunks: list[str]) -> None:
            raise RuntimeError("embedding unavailable")

    service, _, _ = make_service(BrokenIndex())
    info = await service.upload("large.txt", b"x" * 5000, "text/plain")
    assert info.status == "SUCCESS" and info.embed == 0
    assert (await service.load_content(info.file_id, "question")).endswith("x" * 5000)


@pytest.mark.asyncio
async def test_delete_failure_does_not_claim_success_or_remove_metadata() -> None:
    class BrokenObjects(MemoryObjects):
        async def delete(self, name: str) -> None:
            raise RuntimeError("MinIO unavailable")

    service, files, _ = make_service()
    service._objects = BrokenObjects()
    info = await service.upload("note.txt", b"abc", "text/plain")
    with pytest.raises(RuntimeError, match="MinIO unavailable"):
        await service.delete(info.file_id)
    assert info.file_id in files.items


@pytest.mark.asyncio
async def test_file_agent_rejects_model_supplied_different_file_id() -> None:
    service, _, _ = make_service()
    info = await service.upload("small.txt", b"hello", "text/plain")
    sessions = MemorySessions()
    runner = FileChatRunner(ChatModel(), service, sessions, max_rounds=3)
    events = [event async for event in runner.stream("问", "conversation", info.file_id)]
    assert events[-1].content == "文件内容不足，无法回答。"
    assert sessions.saved["file_id"] == info.file_id
    assert sessions.saved["answer"] == "文件内容不足，无法回答。"


def test_legacy_file_routes_shape_without_external_services() -> None:
    service, _, _ = make_service()
    original = app.state.file_service
    app.state.file_service = service
    try:
        with TestClient(app) as client:
            uploaded = client.post(
                "/file/upload", files={"file": ("note.txt", b"hello", "text/plain")}
            ).json()
            file_id = uploaded["data"]["fileId"]
            listed = client.get("/file/list").json()
            content = client.get(f"/file/content/{file_id}").json()
            exists = client.get(f"/file/exists/{file_id}").json()
            deleted = client.delete(f"/file/{file_id}").json()
    finally:
        app.state.file_service = original
    assert uploaded["code"] == 200
    assert listed["data"]["files"][file_id]["fileName"] == "note.txt"
    assert content["data"] == {"content": "hello", "length": 5}
    assert exists["data"] is True
    assert deleted == {"code": 200, "message": "文件删除成功", "data": None}


def test_pgvector_url_uses_java_compatible_fields() -> None:
    """Java 分项配置应生成 psycopg 异步连接串，并正确转义认证字符。"""
    settings = Settings(
        _env_file=None,
        pgvector_host="127.0.0.1",
        pgvector_port=5432,
        pgvector_database="vectordb",
        pgvector_user="post@gres",
        pgvector_password="p:a/ss",
    )
    assert settings.resolved_pgvector_url() == (
        "postgresql+psycopg://post%40gres:p%3Aa%2Fss@127.0.0.1:5432/vectordb"
    )


def test_pgvector_legacy_url_overrides_java_fields() -> None:
    """已有部署继续使用 PGVECTOR_URL 时，应优先采用该显式兼容配置。"""
    settings = Settings(
        _env_file=None,
        pgvector_url="postgresql+psycopg://legacy@db/old",
        pgvector_database="vectordb",
    )
    assert settings.resolved_pgvector_url() == "postgresql+psycopg://legacy@db/old"
